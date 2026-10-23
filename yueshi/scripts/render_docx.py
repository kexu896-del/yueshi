#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
render_docx.py —— plan.json → Word 渲染器（四渲染器之一，只读 plan.json）。

用法：python scripts/render_docx.py --input plan.json --output plan.docx

实现：复用 render_markdown 的章节结构（同一 plan.json、同一门禁），
优先调 pandoc 转 docx；无 pandoc 时用 python-docx 做基础排版兜底。
"""
import argparse, json, os, shutil, subprocess, sys, tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
# round74：门禁统一走 render_gates（与 Markdown 同一套计划级 + 文字门禁）。
from render_gates import enforce, validate_user_text  # noqa: E402
from render_markdown import _md  # noqa: E402


def _md_to_docx_pandoc(md_path, out_path):
    subprocess.run(["pandoc", md_path, "-o", out_path], check=True)


def _md_to_docx_fallback(md_text, out_path):
    """python-docx 兜底：标题/表格/列表的基础排版（无 pandoc 时使用）。"""
    from docx import Document
    doc = Document()
    for line in md_text.splitlines():
        s = line.rstrip()
        if not s:
            continue
        if s.startswith("### "):
            doc.add_heading(s[4:], level=3)
        elif s.startswith("## "):
            doc.add_heading(s[3:], level=2)
        elif s.startswith("# "):
            doc.add_heading(s[2:], level=1)
        elif s.startswith("|---"):
            continue
        elif s.startswith("|") and s.endswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
            doc.add_paragraph(" | ".join(cells))
        elif s.startswith("- "):
            doc.add_paragraph(s[2:], style="List Bullet")
        elif s.startswith("> "):
            doc.add_paragraph(s[2:]).italic = True
        else:
            doc.add_paragraph(s.replace("**", ""))
    doc.save(out_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True)
    ap.add_argument("--output", required=True)
    a = ap.parse_args()
    plan = json.load(open(a.input, encoding="utf-8"))
    # round74：四格式统一门禁（必填键/版本/泄漏/价格/流转），并新增产物文字门禁。
    enforce(plan, "docx")
    md = _md(plan)
    errs = validate_user_text(md)
    if errs:
        print("渲染门禁未通过：" + "; ".join(errs), file=sys.stderr)
        sys.exit(3)
    if shutil.which("pandoc"):
        with tempfile.NamedTemporaryFile("w", suffix=".md", delete=False, encoding="utf-8") as f:
            f.write(md)
            tmp = f.name
        try:
            _md_to_docx_pandoc(tmp, a.output)
        finally:
            os.unlink(tmp)
    else:
        _md_to_docx_fallback(md, a.output)
    print(f"Word 已生成: {a.output}")


if __name__ == "__main__":
    main()
